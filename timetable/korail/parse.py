"""코레일 시각표 엑셀(일반열차·고속열차)을 노선·방향별 표로 정리한다 → JSON
블록은 시트 안에서 가로·세로 어디에 놓여도 앵커 셀('시발역' / '열차번호')로 찾는다."""
import openpyxl, json, re, sys, collections, unicodedata
SIMP = str.maketrans('东庆晋尔', '東慶晉爾')   # 고속열차 엑셀에 섞인 간체자만 고친다 (里·咸 등은 한국 한자 그대로)
NATIVE = {'서울'}          # 한자가 없는 고유어 역명 — 서울은 서울
FIX = {'남창': '南倉', '각계': '覺溪', '군위': '軍威', '진부(오대산)': '珍富(五臺山)'}   # 원본 엑셀의 잘못

def tmin(v):
    """'HH:MM:SS' → 분 (초는 버린다). 0:00:00 은 서지 않음(None). 글자면 그대로"""
    if v is None: return None
    if hasattr(v, 'hour'):
        m = v.hour * 60 + v.minute
        return m if (m or v.second) else None
    s = str(v).strip()
    if not s: return None
    m = re.fullmatch(r'(\d{1,2}):(\d{2})(?::(\d{2}))?', s)
    if m:
        h, mi = int(m[1]), int(m[2])
        if h == 0 and mi == 0 and (m[3] or '00') == '00': return None
        return h * 60 + mi
    return clean(s)

def clean(s): return re.sub(r'\s+', ' ', str(s if s is not None else '').replace('\xa0', ' ')).strip()

def note(v):
    s = clean(v)
    return '' if s in ('매일', '00:00:00', '0:00:00') else s

hanja = {}
def addh(ko, hj):
    ko, hj = clean(ko), unicodedata.normalize('NFKC', clean(hj)).replace(' ', '')   # 호환 한자(金 U+F90A 등)도 정자로
    hj = re.sub(r'[()（）]', '', hj)
    if ko and hj and ko not in NATIVE and re.fullmatch(r'[一-鿿]+', hj):
        hanja.setdefault(ko, FIX.get(ko) or hj.translate(SIMP))

def grid(ws):
    """병합 셀을 펼친 2차원 값 배열"""
    R, C = ws.max_row, ws.max_column
    g = [[None] * (C + 1) for _ in range(R + 1)]
    for row in ws.iter_rows(min_row=1, max_row=R, max_col=C):
        for c in row:
            if c.value is not None: g[c.row - 1][c.column - 1] = c.value
    for mr in ws.merged_cells.ranges:
        v = g[mr.min_row - 1][mr.min_col - 1]
        for r in range(mr.min_row, mr.max_row + 1):
            for cc in range(mr.min_col, mr.max_col + 1): g[r - 1][cc - 1] = v
    return g

def title_dir(g, r, c):
    for k in range(r, max(-1, r - 8), -1):
        for cc in (c, c - 1, c + 1):
            if 0 <= cc < len(g[k]):
                s = clean(g[k][cc])
                if re.search(r'하행|下行', s): return 'down'
                if re.search(r'상행|上行', s): return 'up'
    return None

def gen_book(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    tables = []
    for ws in wb.worksheets:
        if ws.title == '보는방법': continue
        g = grid(ws)
        for r0, row in enumerate(g):
            for c, v in enumerate(row):
                if clean(v) != '시발역' or (r0 and clean(g[r0 - 1][c]) == '시발역'): continue
                col = lambda r: clean(g[r][c])
                kind_r = next(k for k in range(r0, r0 + 8) if col(k) == '열차종별')
                no_r = kind_r + 1
                end = next(k for k in range(no_r + 1, len(g)) if col(k) in ('비 고', '비고', '종착역'))
                note_r = end if col(end) in ('비 고', '비고') else None
                term_r = end + 1 if note_r is not None else end
                # 열차 열: 시발역 오른쪽, '始發驛' 같은 한자 라벨이 나올 때까지
                cols, hc = [], None
                for j in range(c + 1, len(g[r0])):
                    s = clean(g[r0][j])
                    if re.fullmatch(r'[一-鿿]+', s) and not clean(g[no_r][j]).isdigit(): hc = j; break
                    if clean(g[no_r][j]): cols.append(j)
                    elif not s: break
                st_rows = [k for k in range(no_r + 1, end) if col(k)]
                stations = [col(k) for k in st_rows]
                if hc is not None:
                    for k in st_rows: addh(col(k), g[k][hc])
                trains = []
                for j in cols:
                    addh(g[r0][j], g[r0 + 1][j]); addh(g[term_r][j], g[term_r + 1][j])
                    trains.append({
                        'no': clean(g[no_r][j]), 'type': clean(g[kind_r][j]),
                        'from': clean(g[r0][j]), 'fromT': tmin(g[r0 + 3][j]),
                        'to': clean(g[term_r][j]), 'toT': tmin(g[term_r + 3][j]),
                        'note': note(g[note_r][j]) if note_r is not None else '',
                        'times': [tmin(g[k][j]) for k in st_rows],
                    })
                tables.append({'line': ws.title, 'cat': 'gen', 'dir': title_dir(g, r0, c),
                               'stations': stations, 'trains': trains})
    return tables

def ktx_book(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    tables = []
    for ws in wb.worksheets:
        g = grid(ws)
        for i, row in enumerate(g):
            for c, v in enumerate(row):
                if clean(v) != '열차번호' or clean(g[i - 1][c]) == '열차번호': continue
                stc, notec = [], None
                for j in range(c + 2, len(row)):
                    s = clean(row[j])
                    if s.startswith('비고'): notec = j; break
                    if not s: break
                    stc.append(j)
                stations = [clean(row[j]) for j in stc]
                for j in stc: addh(row[j], g[i + 1][j])
                trains, k = [], i + 3
                while k < len(g) and clean(g[k][c]):
                    rr = g[k]
                    trains.append({'no': clean(rr[c]), 'type': clean(rr[c + 1]),
                                   'note': note(rr[notec]) if notec else '',
                                   'times': [tmin(rr[j]) for j in stc]})
                    k += 1
                tables.append({'line': ws.title, 'cat': 'ktx', 'dir': title_dir(g, i, c),
                               'stations': stations, 'trains': trains})
    return tables

if __name__ == '__main__':
    gen, ktx = sys.argv[1], sys.argv[2]
    G = gen_book(gen); T = ktx_book(ktx) + G   # 한자는 일반열차 쪽(정자체)을 먼저
    merged = collections.OrderedDict()
    for t in T:
        k = (t['line'], t['cat'], t['dir'], tuple(t['stations']))
        if k in merged: merged[k]['trains'] += t['trains']
        else: merged[k] = t
    out = {'tables': list(merged.values()), 'hanja': hanja}
    json.dump(out, open(sys.argv[3], 'w'), ensure_ascii=False, indent=0)
    for t in out['tables']: print(t['cat'], t['line'], t['dir'], len(t['stations']), 'st', len(t['trains']), 'trains')
    print(len(hanja), 'hanja names')
